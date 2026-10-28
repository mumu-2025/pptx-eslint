"""speaker-notes / notes-for-charts — notes presence for async & chart a11y."""
from __future__ import annotations

from typing import Any

from pptx.enum.shapes import MSO_SHAPE_TYPE

from ..models import Violation


def _notes_text(slide) -> str:
    try:
        if not slide.has_notes_slide:
            return ""
        tf = slide.notes_slide.notes_text_frame
        return (tf.text or "").strip() if tf is not None else ""
    except Exception:
        return ""


def _has_chart(slide) -> bool:
    for shape in slide.shapes:
        try:
            if shape.shape_type == MSO_SHAPE_TYPE.CHART:
                return True
            if getattr(shape, "has_chart", False):
                return True
        except Exception:
            continue
    return False


def _body_char_count(slide) -> int:
    n = 0
    for shape in slide.shapes:
        if not getattr(shape, "has_text_frame", False):
            continue
        for p in shape.text_frame.paragraphs:
            n += len((p.text or "").strip())
    return n


def check_speaker_notes(slide, idx: int, rule: dict[str, Any], ctx: dict[str, Any]) -> list[Violation]:
    """Require notes when slide has enough body text (async-doc / reading handoff)."""
    min_chars = int(rule.get("min_body_chars", 40))
    skip_first = bool(rule.get("skip_first_slide", True))
    if skip_first and idx == 1:
        return []
    body = _body_char_count(slide)
    if body < min_chars:
        return []
    notes = _notes_text(slide)
    min_notes = int(rule.get("min_notes_chars", 20))
    if len(notes) >= min_notes:
        return []
    return [
        Violation(
            rule_id="speaker-notes",
            severity=rule.get("_severity", "warning"),
            slide=idx,
            message_en=f"content slide ({body} body chars) lacks speaker notes (≥{min_notes} chars)",
            message_zh=f"内容页（正文 {body} 字符）缺少演讲者备注（需 ≥{min_notes} 字符）",
            fix_hint_en="Add speaker notes so async readers get the spoken narrative.",
            fix_hint_zh="补充演讲者备注，让异步读者也能看到口述叙事。",
            category="design",
        )
    ]


def check_notes_for_charts(slide, idx: int, rule: dict[str, Any], ctx: dict[str, Any]) -> list[Violation]:
    if not _has_chart(slide):
        return []
    notes = _notes_text(slide)
    min_notes = int(rule.get("min_notes_chars", 30))
    if len(notes) >= min_notes:
        return []
    return [
        Violation(
            rule_id="notes-for-charts",
            severity=rule.get("_severity", "warning"),
            slide=idx,
            message_en=f"chart slide missing explanatory notes (≥{min_notes} chars)",
            message_zh=f"含图表页缺少说明性备注（需 ≥{min_notes} 字符）",
            fix_hint_en="Summarize the chart takeaway in speaker notes (a11y + async reading).",
            fix_hint_zh="在备注中概括图表结论（无障碍与异步阅读）。",
            category="chart",
        )
    ]
