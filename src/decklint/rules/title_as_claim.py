"""title-as-claim — ban vague titles (bilingual list)."""
from __future__ import annotations

from typing import Any

from pptx.util import Pt

from ..models import Violation

DEFAULT_BAN = [
    "Overview",
    "Introduction",
    "Agenda",
    "Summary",
    "Conclusion",
    "Thank you",
    "Thanks",
    "Q&A",
    "QA",
    "Contents",
    "Table of Contents",
    "Background",
    "Update",
    "Misc",
    "概况",
    "概述",
    "简介",
    "引言",
    "目录",
    "议程",
    "总结",
    "结论",
    "谢谢",
    "致谢",
    "背景",
    "汇报",
    "分享",
]


def check(slide, idx: int, rule: dict[str, Any], ctx: dict[str, Any]) -> list[Violation]:
    ban = [b.lower() for b in (rule.get("ban_phrases") or DEFAULT_BAN)]
    candidates: list[str] = []
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        full = "\n".join(p.text for p in shape.text_frame.paragraphs).strip()
        if not full or len(full) > 80:
            continue
        try:
            if shape.top is not None and shape.top < Pt(160):
                candidates.append(full)
        except Exception:
            candidates.append(full)
    # also first text box
    if not candidates:
        for shape in slide.shapes:
            if shape.has_text_frame:
                full = "\n".join(p.text for p in shape.text_frame.paragraphs).strip()
                if full:
                    candidates.append(full.split("\n")[0])
                    break

    out: list[Violation] = []
    for t in candidates[:3]:
        low = t.lower().strip()
        for b in ban:
            if low == b or low.startswith(b + " ") or low.startswith(b + "：") or low.startswith(b + ":"):
                out.append(
                    Violation(
                        rule_id="title-as-claim",
                        severity=rule.get("_severity", "warning"),
                        slide=idx,
                        message_en=f"vague title «{t[:40]}» — prefer a claim",
                        message_zh=f"空泛标题「{t[:40]}」— 请改成主张句",
                        fix_hint_en="State the takeaway as a full sentence claim.",
                        fix_hint_zh="把结论写成完整主张句。",
                        category="design",
                    )
                )
                break
    return out
