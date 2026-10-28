"""text-contrast — WCAG ratio when solid RGB resolvable."""
from __future__ import annotations

from typing import Any

from ..contrast import contrast_ratio, is_large_text, passes_aa
from ..models import Violation
from ._shapes import (
    iter_text_runs,
    run_font_bold,
    run_font_rgb,
    run_font_size_pt,
    shape_solid_fill_rgb,
)


def check(slide, idx: int, rule: dict[str, Any], ctx: dict[str, Any]) -> list[Violation]:
    # slide background solid if available
    slide_bg = _slide_bg_rgb(slide)
    # collect opaque filled shapes as possible backgrounds (back to front)
    filled: list[tuple[tuple[float, float, float, float], tuple[int, int, int]]] = []
    from ._shapes import shape_box_pt

    for shape in slide.shapes:
        rgb = shape_solid_fill_rgb(shape)
        box = shape_box_pt(shape)
        if rgb and box:
            filled.append((box, rgb))

    out: list[Violation] = []
    min_body = float(rule.get("min_ratio", 4.5))
    min_large = float(rule.get("min_ratio_large", 3.0))
    seen = 0

    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        shape_fill = shape_solid_fill_rgb(shape)
        tbox = shape_box_pt(shape)
        for _p, run, text in iter_text_runs(shape):
            if not text.strip():
                continue
            fg = run_font_rgb(run)
            if not fg:
                continue
            bg = shape_fill
            if bg is None and tbox is not None:
                # find topmost filled shape that fully contains text box center
                cx = (tbox[0] + tbox[2]) / 2
                cy = (tbox[1] + tbox[3]) / 2
                for box, rgb in reversed(filled):
                    if box[0] <= cx <= box[2] and box[1] <= cy <= box[3]:
                        bg = rgb
                        break
            if bg is None:
                bg = slide_bg
            if bg is None:
                continue
            ratio = contrast_ratio(fg, bg)
            pt = run_font_size_pt(run)
            bold = run_font_bold(run)
            large = is_large_text(pt, bold)
            threshold = min_large if large else min_body
            ok = passes_aa(ratio, large) if rule.get("use_wcag_aa", True) else ratio >= threshold
            # also respect explicit thresholds
            if ratio < threshold:
                ok = False
            if ok:
                continue
            seen += 1
            if seen > 5:
                break
            out.append(
                Violation(
                    rule_id="text-contrast",
                    severity=rule.get("_severity", "error"),
                    slide=idx,
                    message_en=(
                        f"contrast {ratio:.2f}:1 < {threshold:.1f}:1 "
                        f"(fg={_hex(fg)} bg={_hex(bg)}, {'large' if large else 'body'})"
                    ),
                    message_zh=(
                        f"对比度 {ratio:.2f}:1 < {threshold:.1f}:1 "
                        f"（前景={_hex(fg)} 背景={_hex(bg)}，{'大字' if large else '正文'}）"
                    ),
                    fix_hint_en="Darken/lighten text or fill until WCAG AA passes.",
                    fix_hint_zh="加深或减淡文字/填充，直到满足 WCAG AA。",
                    category="a11y",
                )
            )
        if seen > 5:
            break
    return out


def _slide_bg_rgb(slide) -> tuple[int, int, int] | None:
    try:
        fill = slide.background.fill
        from pptx.enum.dml import MSO_FILL

        if fill.type == MSO_FILL.SOLID:
            rgb = fill.fore_color.rgb
            return (int(rgb[0]), int(rgb[1]), int(rgb[2]))
    except Exception:
        return None
    return None


def _hex(rgb: tuple[int, int, int]) -> str:
    return f"#{rgb[0]:02X}{rgb[1]:02X}{rgb[2]:02X}"
