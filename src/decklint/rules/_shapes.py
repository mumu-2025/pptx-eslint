"""Shared shape / color helpers for rules."""
from __future__ import annotations

from typing import Any, Iterator

from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.util import Emu, Pt


def emu_to_pt(emu: int | None) -> float | None:
    if emu is None:
        return None
    return Emu(emu).pt


def shape_box_pt(shape) -> tuple[float, float, float, float] | None:
    """Return (left, top, right, bottom) in points, or None."""
    try:
        if shape.left is None or shape.top is None:
            return None
        if shape.width is None or shape.height is None:
            return None
        l = shape.left.pt
        t = shape.top.pt
        w = shape.width.pt
        h = shape.height.pt
        return (l, t, l + w, t + h)
    except Exception:
        return None


def overlap_area(a: tuple[float, float, float, float], b: tuple[float, float, float, float]) -> float:
    x1 = max(a[0], b[0])
    y1 = max(a[1], b[1])
    x2 = min(a[2], b[2])
    y2 = min(a[3], b[3])
    if x2 <= x1 or y2 <= y1:
        return 0.0
    return (x2 - x1) * (y2 - y1)


def box_area(b: tuple[float, float, float, float]) -> float:
    return max(0.0, (b[2] - b[0]) * (b[3] - b[1]))


def iter_text_runs(shape) -> Iterator[tuple[Any, Any, str]]:
    """Yield (paragraph, run, text) for shapes with text frames."""
    if not getattr(shape, "has_text_frame", False):
        return
    for p in shape.text_frame.paragraphs:
        for run in p.runs:
            yield p, run, run.text or ""


def solid_rgb_from_color(color) -> tuple[int, int, int] | None:
    """Best-effort extract RGB from a ColorFormat-like object."""
    try:
        rgb = color.rgb
        if rgb is None:
            return None
        return (int(rgb[0]), int(rgb[1]), int(rgb[2]))
    except Exception:
        return None


def shape_solid_fill_rgb(shape) -> tuple[int, int, int] | None:
    try:
        fill = shape.fill
        # solid fill
        from pptx.enum.dml import MSO_FILL

        if fill.type != MSO_FILL.SOLID:
            return None
        return solid_rgb_from_color(fill.fore_color)
    except Exception:
        return None


def run_font_rgb(run) -> tuple[int, int, int] | None:
    try:
        return solid_rgb_from_color(run.font.color)
    except Exception:
        return None


def run_font_size_pt(run) -> float | None:
    try:
        if run.font.size is None:
            return None
        return float(run.font.size.pt)
    except Exception:
        return None


def run_font_bold(run) -> bool | None:
    try:
        return bool(run.font.bold) if run.font.bold is not None else None
    except Exception:
        return None


def run_font_name(run) -> str | None:
    try:
        name = run.font.name
        return str(name) if name else None
    except Exception:
        return None


def is_picture(shape) -> bool:
    try:
        return shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    except Exception:
        return False


def picture_descr(shape) -> str | None:
    """Read cNvPr descr / title if present on picture element."""
    try:
        el = shape._element  # noqa: SLF001
        # look for cNvPr
        for node in el.iter():
            tag = node.tag.split("}")[-1] if "}" in node.tag else node.tag
            if tag == "cNvPr":
                descr = node.get("descr") or node.get("title") or ""
                return descr.strip() or None
    except Exception:
        return None
    return None


def slide_has_title(slide) -> bool:
    try:
        shapes = slide.shapes
        if getattr(shapes, "title", None) is not None and shapes.title is not None:
            text = shapes.title.text_frame.text if shapes.title.has_text_frame else ""
            return bool(text.strip())
    except Exception:
        pass
    return False


EMU_PER_INCH = 914400


def slide_size_inches(prs) -> tuple[float, float]:
    return (prs.slide_width / EMU_PER_INCH, prs.slide_height / EMU_PER_INCH)
