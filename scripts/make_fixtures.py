#!/usr/bin/env python3
"""Build bad/good sample decks for local demos and CI."""
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1] / "fixtures"


def bad() -> None:
    out = ROOT / "bad"
    out.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12), Inches(6.5))
    tf = box.text_frame
    tf.word_wrap = True
    run = tf.paragraphs[0].add_run()
    run.text = "Overview"
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(180, 180, 180)
    for i in range(12):
        p = tf.add_paragraph()
        p.text = (
            f"- This is a very long bullet point number {i} that tries to put an "
            "entire paragraph of document text onto a projected slide which nobody "
            "can read from the back of the room"
        )
        for r in p.runs:
            r.font.size = Pt(12)
    prs.save(out / "slideument.pptx")


def good() -> None:
    out = ROOT / "good"
    out.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(10, 10, 12)
    bg.line.fill.background()
    title = s.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(11.5), Inches(2))
    run = title.text_frame.paragraphs[0].add_run()
    run.text = "Retention rose 17% after we cut onboarding to three steps"
    run.font.size = Pt(36)
    run.font.bold = True
    run.font.color.rgb = RGBColor(255, 255, 255)
    prs.save(out / "claim-slide.pptx")


if __name__ == "__main__":
    bad()
    good()
    print("wrote fixtures/bad/slideument.pptx and fixtures/good/claim-slide.pptx")
